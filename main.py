
# -*- coding: utf-8 -*-
import sys
import os
import re
import subprocess
import unicodedata
import zipfile
import xml.etree.ElementTree as ET
import mimetypes
from html import unescape

from PySide6.QtCore import Qt
from PySide6.QtGui import QIcon
from PySide6.QtWidgets import (
    QApplication,
    QWidget,
    QLabel,
    QPushButton,
    QVBoxLayout,
    QFileDialog,
    QListWidget,
    QListWidgetItem,
    QTableWidget,
    QTabWidget,
    QLineEdit,
    QTableWidgetItem,
    QDialog,
    QTextEdit
    ,
)

from database.database import PasswordVault, PasswordVaultError



app = QApplication(sys.argv)

password_vault = PasswordVault()
password_cache = {}

def resource_path(relative_path):
    base = getattr(sys, "_MEIPASS", os.path.dirname(os.path.abspath(__file__)))
    return os.path.join(base, relative_path)

current_language = "en_uk"
last_scan_result = None
protection_needs_action = False
home_firewall_enabled = False


TEXTS = {
    "en_uk": {
        "status_secure": "Status: Ready to scan",
        "status_scanning": "Status: Scanning...",
        "status_scan_completed":
            "Status: Scan completed - inspected {scanned} files, "
            "found {found} flagged items",
        "quick_scan": "Quick Scan (Downloads)",
        "custom_scan": "Custom Scan",
        "choose_folder": "Choose folder for custom scanning",
        "scan_failed": "Status: Scan failed - folder unavailable",
        "suspicious_extensions":
            "Executable/script file requiring review: {path}",
        "document_flagged": "Profanity detected in document: {path}",
        "document_clean": "No configured profanity detected: {path}",
        "document_empty":
            "No extractable text found; file may be scanned, "
            "protected or unsupported: {path}",
        "file_unavailable":
            "File identified, but its contents are not text-previewable: {path}",
        "no_documents_found": "No supported documents were found.",
        "home_secure": "Protection status: No flagged items in last scan",
        "home_action_needed": "Protection status: Review flagged items",
        "home_scan_default": "Last scan: Not yet scanned.",
        "home_scan_result":
            "Last scan: inspected {scanned} files, found {found} flagged items",
        "home_firewall_default": "Firewall: Not yet configured",
        "home_firewall_enabled": "Firewall: Enabled",
        "firewall_title": "Firewall & Network Protection",
        "domain_unknown": "Domain network: Firewall status not checked",
        "private_unknown": "Private network: Firewall status not checked",
        "public_unknown": "Public network: Firewall status not checked",
        "domain_button": "Turn on Domain network firewall",
        "private_button": "Turn on Private network firewall",
        "public_button": "Turn on Public network firewall",
        "domain_name": "Domain network",
        "private_name": "Private network",
        "public_name": "Public network",
        "firewall_on": "{name}: Firewall is on.",
        "firewall_enabled": "Firewall status: {name} enabled",
        "firewall_failed":
            "Firewall status: Failed - try running Viverium as administrator",
        "password_locked": "Password manager: Locked",
        "password_create_master":
            "Password manager: Create a master password",
        "password_master_placeholder": "Master password",
        "password_unlock_button": "Unlock / create master password",
        "password_service_placeholder": "Service",
        "password_username_placeholder": "Username",
        "password_password_placeholder": "Password",
        "password_save_button": "Save password",
        "password_table_id": "ID",
        "password_table_service": "Service",
        "password_table_username": "Username",
        "password_table_password": "Password",
        "password_show_button": "Show selected password",
        "password_delete_button": "Delete selected password",
        "selected_hidden": "Selected password: Hidden",
        "selected_select_first":
            "Selected password: Select an item first",
        "selected_not_found": "Not found",
        "selected_value": "Selected password: {password}",
        "password_unlocked": "Password manager: Unlocked",
        "password_master_created":
            "Password manager: Master password created",
        "password_saved": "Password manager: Password saved",
        "password_select_first":
            "Password manager: Select an item first",
        "password_deleted": "Password manager: Password deleted",
        "password_error": "Password manager: {error}",
        "language_status": "Language: English (UK)",
        "language_description": "Choose the programme language",
        "language_ptbr": "Português (Brasil)",
        "language_enuk": "English (UK)",
        "tab_home": "Home",
        "tab_scan": "Scanning",
        "tab_firewall": "Firewall & Network protection",
        "tab_password": "Password manager",
        "tab_language": "Language",
        "preview_title": "File preview: {name}",
        "preview_unavailable":
            "This format cannot be displayed as plain text.",
        "preview_error": "Could not read this file:\n{error}",
        "preview_hint":
            "Double-click a result to inspect the file content or details.",
        "close": "Close",
    },

    "pt_br": {
        "status_secure": "Status: Pronto para verificar",
        "status_scanning": "Status: Verificando...",
        "status_scan_completed":
            "Status: Verificação concluída - {scanned} arquivos "
            "inspecionados, {found} itens sinalizados",
        "quick_scan": "Verificação rápida (Downloads)",
        "custom_scan": "Verificação personalizada",
        "choose_folder": "Escolha a pasta para verificar",
        "scan_failed": "Status: Falha - pasta indisponível",
        "suspicious_extensions":
            "Arquivo executável/script que precisa ser analisado: {path}",
        "document_flagged": "Palavrão detectado no documento: {path}",
        "document_clean": "Nenhum palavrão configurado detectado: {path}",
        "document_empty":
            "Nenhum texto extraível; o arquivo pode ser digitalizado, "
            "protegido ou incompatível: {path}",
        "file_unavailable":
            "Arquivo identificado, mas o conteúdo não pode ser exibido "
            "como texto: {path}",
        "no_documents_found":
            "Nenhum documento compatível foi encontrado.",
        "home_secure":
            "Proteção: Nenhum item sinalizado na última verificação",
        "home_action_needed": "Proteção: Revise os itens sinalizados",
        "home_scan_default": "Última verificação: Ainda não realizada.",
        "home_scan_result":
            "Última verificação: {scanned} arquivos inspecionados, "
            "{found} itens sinalizados",
        "home_firewall_default": "Firewall: Ainda não configurado",
        "home_firewall_enabled": "Firewall: Ativado",
        "firewall_title": "Firewall e proteção de rede",
        "domain_unknown": "Rede de domínio: Firewall não verificado",
        "private_unknown": "Rede privada: Firewall não verificado",
        "public_unknown": "Rede pública: Firewall não verificado",
        "domain_button": "Ativar firewall da rede de domínio",
        "private_button": "Ativar firewall da rede privada",
        "public_button": "Ativar firewall da rede pública",
        "domain_name": "Rede de domínio",
        "private_name": "Rede privada",
        "public_name": "Rede pública",
        "firewall_on": "{name}: Firewall ativado.",
        "firewall_enabled": "Status do firewall: {name} ativado",
        "firewall_failed":
            "Firewall: Falha - tente executar como administrador",
        "password_locked": "Gerenciador de senhas: Bloqueado",
        "password_create_master":
            "Gerenciador de senhas: Crie uma senha mestre",
        "password_master_placeholder": "Senha mestre",
        "password_unlock_button": "Desbloquear / criar senha mestre",
        "password_service_placeholder": "Serviço",
        "password_username_placeholder": "Usuário",
        "password_password_placeholder": "Senha",
        "password_save_button": "Salvar senha",
        "password_table_id": "ID",
        "password_table_service": "Serviço",
        "password_table_username": "Usuário",
        "password_table_password": "Senha",
        "password_show_button": "Mostrar senha selecionada",
        "password_delete_button": "Apagar senha selecionada",
        "selected_hidden": "Senha selecionada: Oculta",
        "selected_select_first":
            "Senha selecionada: Selecione um item primeiro",
        "selected_not_found": "Não encontrada",
        "selected_value": "Senha selecionada: {password}",
        "password_unlocked": "Gerenciador de senhas: Desbloqueado",
        "password_master_created":
            "Gerenciador de senhas: Senha mestre criada",
        "password_saved": "Gerenciador de senhas: Senha salva",
        "password_select_first":
            "Gerenciador de senhas: Selecione um item primeiro",
        "password_deleted": "Gerenciador de senhas: Senha apagada",
        "password_error": "Gerenciador de senhas: {error}",
        "language_status": "Idioma: Português (Brasil)",
        "language_description": "Escolha o idioma do programa",
        "language_ptbr": "Português (Brasil)",
        "language_enuk": "English (UK)",
        "tab_home": "Início",
        "tab_scan": "Verificação",
        "tab_firewall": "Firewall e proteção de rede",
        "tab_password": "Gerenciador de senhas",
        "tab_language": "Idioma",
        "preview_title": "Pré-visualização: {name}",
        "preview_unavailable":
            "Este formato não pode ser exibido como texto simples.",
        "preview_error": "Não foi possível ler este arquivo:\n{error}",
        "preview_hint":
            "Dê dois cliques em um resultado para visualizar o conteúdo.",
        "close": "Fechar",
    },
}


def tr(key, **values):
    return TEXTS[current_language][key].format(**values)



TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".tsv", ".log", ".json", ".xml",
    ".html", ".htm", ".rtf", ".docx", ".pdf", ".xlsx", ".xlsm",
    ".pptx", ".odt", ".ini", ".yaml", ".yml", ".py", ".js",
    ".css", ".sql", ".c", ".cpp", ".h", ".java", ".bat", ".ps1",
    ".sh", ".conf", ".toml",
}

SUSPICIOUS_EXTENSIONS = (
    ".exe", ".dll", ".msi", ".bat", ".cmd", ".vbs", ".ps1", ".scr"
)

MAX_TEXT_BYTES = 8 * 1024 * 1024

BAD_WORDS = [
    # Português
    "caralho", "porra", "puta", "merda", "foda", "bosta",
    "cacete", "pqp", "filho da puta",

    # Inglês
    "fuck", "fucking", "shit", "bullshit", "bitch", "bastard",
    "motherfucker", "asshole", "dick", "cunt",
]


def normalize_text(text):
    text = unicodedata.normalize("NFKD", text)
    text = text.encode("ascii", "ignore").decode("ascii")
    return text.lower()



def xml_text(xml_bytes):
    root = ET.fromstring(xml_bytes)
    pieces = []

    for element in root.iter():
        if element.text and element.text.strip():
            pieces.append(element.text.strip())

    return "\n".join(pieces)


def read_docx_text(file_path):
    pieces = []

    with zipfile.ZipFile(file_path, "r") as archive:
        for name in archive.namelist():
            if name.startswith("word/") and name.endswith(".xml"):
                try:
                    pieces.append(xml_text(archive.read(name)))
                except ET.ParseError:
                    continue

    return "\n".join(pieces)


def read_pdf_text(file_path):
    from pypdf import PdfReader

    reader = PdfReader(file_path)
    pieces = []

    for page in reader.pages:
        pieces.append(page.extract_text() or "")

    return "\n\n".join(pieces)


def read_xlsx_text(file_path):
    from openpyxl import load_workbook

    workbook = load_workbook(
        file_path,
        read_only=True,
        data_only=True,
    )

    pieces = []

    try:
        for sheet in workbook.worksheets:
            pieces.append(f"[Planilha: {sheet.title}]")

            for row in sheet.iter_rows(values_only=True):
                values = [
                    str(value)
                    for value in row
                    if value is not None
                ]

                if values:
                    pieces.append(" | ".join(values))
    finally:
        workbook.close()

    return "\n".join(pieces)


def read_pptx_text(file_path):
    from pptx import Presentation

    presentation = Presentation(file_path)
    pieces = []

    for index, slide in enumerate(presentation.slides, start=1):
        pieces.append(f"[Slide {index}]")

        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                if shape.text.strip():
                    pieces.append(shape.text)

            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    pieces.append(
                        " | ".join(cell.text for cell in row.cells)
                    )

    return "\n".join(pieces)


def read_odt_text(file_path):
    with zipfile.ZipFile(file_path, "r") as archive:
        return xml_text(archive.read("content.xml"))


def clean_rtf(text):
    text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
    text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
    text = text.replace("{", " ").replace("}", " ")
    return re.sub(r"\s+", " ", text).strip()


def read_document_text(file_path):
    extension = os.path.splitext(file_path)[1].lower()
    file_size = os.path.getsize(file_path)

    if (
        file_size > MAX_TEXT_BYTES
        and extension not in {
            ".pdf", ".docx", ".xlsx", ".xlsm", ".pptx", ".odt"
        }
    ):
        raise ValueError(
            "Arquivo muito grande para pré-visualização automática."
        )

    if extension == ".pdf":
        return read_pdf_text(file_path)

    if extension == ".docx":
        return read_docx_text(file_path)

    if extension in {".xlsx", ".xlsm"}:
        return read_xlsx_text(file_path)

    if extension == ".pptx":
        return read_pptx_text(file_path)

    if extension == ".odt":
        return read_odt_text(file_path)

    with open(file_path, "rb") as file:
        raw = file.read(MAX_TEXT_BYTES + 1)

    if len(raw) > MAX_TEXT_BYTES:
        raw = raw[:MAX_TEXT_BYTES]

    decoded = None

    for encoding in ("utf-8-sig", "utf-16", "cp1252"):
        try:
            decoded = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue

    if decoded is None:
        decoded = raw.decode("utf-8", errors="replace")

    if extension in {".html", ".htm"}:
        decoded = re.sub(
            r"(?is)<(script|style).*?>.*?</\1>",
            " ",
            decoded,
        )
        decoded = re.sub(r"(?s)<[^>]+>", " ", decoded)
        decoded = unescape(decoded)

    elif extension == ".rtf":
        decoded = clean_rtf(decoded)

    return decoded.replace("\x00", " ")


def detect_bad_words(text):
    normalised = normalize_text(text)

    for word in BAD_WORDS:
        normalised_word = normalize_text(word)

        pattern = (
            r"(?<![a-z0-9_])"
            + re.escape(normalised_word)
            + r"(?![a-z0-9_])"
        )

        if re.search(pattern, normalised):
            return True

    return False


window = QWidget()
window.setWindowTitle("Viverium")
window.resize(760, 520)

window.setWindowIcon(QIcon(resource_path("assets/viverium_icon.ico")))

title_label = QLabel("Viverium")
status_label = QLabel(tr("status_secure"))

scan_button = QPushButton(tr("quick_scan"))
custom_scan_button = QPushButton(tr("custom_scan"))

result_list = QListWidget()


def add_result(message, file_path=None):
    item = QListWidgetItem(message)

    if file_path:
        item.setData(Qt.UserRole, file_path)

    result_list.addItem(item)



def preview_file(file_path):
    dialog = QDialog(window)
    dialog.setWindowTitle(
        tr("preview_title", name=os.path.basename(file_path))
    )
    dialog.resize(760, 560)

    layout = QVBoxLayout(dialog)
    preview = QTextEdit()
    preview.setReadOnly(True)

    extension = os.path.splitext(file_path)[1].lower()

    try:
        if extension in TEXT_EXTENSIONS:
            content = read_document_text(file_path)

            if content.strip():
                preview.setPlainText(content)
            else:
                preview.setPlainText(
                    tr("document_empty")
                    + "\n\n"
                    + file_path
                    + "\n\nPDFs digitalizados podem precisar de OCR. "
                    "Documentos protegidos por senha podem não ser lidos."
                )

        elif extension in {
            ".png", ".jpg", ".jpeg", ".bmp", ".gif", ".webp"
        }:
            preview.setPlainText(
                "Imagem identificada.\n\n"
                "Esta versão ainda não extrai texto de imagens. "
                "Para isso, será necessário integrar OCR.\n\n"
                + file_path
            )

        else:
            mime = mimetypes.guess_type(file_path)[0] or "Tipo desconhecido"
            size = os.path.getsize(file_path)

            preview.setPlainText(
                tr("preview_unavailable")
                + f"\n\nNome: {os.path.basename(file_path)}"
                + f"\nExtensão: {extension or '(sem extensão)'}"
                + f"\nTipo MIME: {mime}"
                + f"\nTamanho: {size:,} bytes"
                + f"\nCaminho: {file_path}"
                + "\n\nExecutáveis não devem ser tratados como documentos "
                "de texto. Esta versão não executa arquivos nem realiza "
                "análise estática completa."
            )

    except Exception as error:
        preview.setPlainText(
            tr("preview_error", error=str(error))
            + "\n\nArquivo: "
            + file_path
        )

    layout.addWidget(preview)

    close_button = QPushButton(tr("close"))
    close_button.clicked.connect(dialog.accept)
    layout.addWidget(close_button)

    dialog.exec()


def open_selected_result(item):
    file_path = item.data(Qt.UserRole)

    if file_path and os.path.isfile(file_path):
        preview_file(file_path)


result_list.itemDoubleClicked.connect(open_selected_result)




def start_scan(scan_folder=None):
    global last_scan_result
    global protection_needs_action

    if scan_folder is None:
        scan_folder = os.path.join(
            os.path.expanduser("~"),
            "Downloads",
        )

    result_list.clear()
    status_label.setText(tr("status_scanning"))
    app.processEvents()

    if not os.path.isdir(scan_folder):
        add_result(tr("scan_failed"))
        status_label.setText(tr("scan_failed"))
        return

    found_count = 0
    scanned_count = 0
    supported_documents = 0

    for folder, _, files in os.walk(scan_folder):
        for filename in files:
            full_path = os.path.join(folder, filename)
            extension = os.path.splitext(filename)[1].lower()

            scanned_count += 1
            flagged = False

            if extension in SUSPICIOUS_EXTENSIONS:
                add_result(
                    tr("suspicious_extensions", path=full_path),
                    full_path,
                )
                found_count += 1
                flagged = True

            if extension in TEXT_EXTENSIONS:
                supported_documents += 1

                try:
                    content = read_document_text(full_path)

                    if not content.strip():
                        add_result(
                            tr("document_empty", path=full_path),
                            full_path,
                        )

                    elif detect_bad_words(content):
                        add_result(
                            tr("document_flagged", path=full_path),
                            full_path,
                        )

                        if not flagged:
                            found_count += 1
                            flagged = True

                    else:
                        add_result(
                            tr("document_clean", path=full_path),
                            full_path,
                        )

                except Exception as error:
                    add_result(
                        f"Falha ao ler / Could not read: "
                        f"{full_path} — {error}",
                        full_path,
                    )

                    print(
                        "Could not inspect file:",
                        full_path,
                        repr(error),
                    )

            elif extension not in SUSPICIOUS_EXTENSIONS:
                add_result(
                    tr("file_unavailable", path=full_path),
                    full_path,
                )

    if supported_documents == 0:
        add_result(tr("no_documents_found"))

    status_label.setText(
        tr(
            "status_scan_completed",
            scanned=scanned_count,
            found=found_count,
        )
    )

    home_scan.setText(
        tr(
            "home_scan_result",
            scanned=scanned_count,
            found=found_count,
        )
    )

    last_scan_result = (scanned_count, found_count)
    protection_needs_action = found_count > 0

    if found_count > 0:
        home_status.setText(tr("home_action_needed"))
    else:
        home_status.setText(tr("home_secure"))


def custom_scan():
    folder = QFileDialog.getExistingDirectory(
        window,
        tr("choose_folder"),
    )

    if folder:
        start_scan(folder)


scan_button.clicked.connect(lambda: start_scan())
custom_scan_button.clicked.connect(custom_scan)


# FIREWALL


tabs = QTabWidget()

home_tab = QWidget()
scan_tab = QWidget()
firewall_tab = QWidget()
password_tab = QWidget()
language_tab = QWidget()

firewall_status = QLabel(tr("firewall_title"))

domain_status = QLabel(tr("domain_unknown"))
private_status = QLabel(tr("private_unknown"))
public_status = QLabel(tr("public_unknown"))

domain_button = QPushButton(tr("domain_button"))
private_button = QPushButton(tr("private_button"))
public_button = QPushButton(tr("public_button"))


def turn_on_firewall_profile(profile, label, display_name):
    global home_firewall_enabled

    try:
        result = subprocess.run(
            [
                "netsh",
                "advfirewall",
                "set",
                profile,
                "state",
                "on",
            ],
            capture_output=True,
            text=True,
            check=False,
        )

        if result.returncode == 0:
            label.setText(
                tr("firewall_on", name=display_name)
            )

            firewall_status.setText(
                tr("firewall_enabled", name=display_name)
            )

            home_firewall.setText(
                tr("home_firewall_enabled")
            )

            home_firewall_enabled = True

        else:
            firewall_status.setText(tr("firewall_failed"))

    except OSError:
        firewall_status.setText(tr("firewall_failed"))


domain_button.clicked.connect(
    lambda: turn_on_firewall_profile(
        "domainprofile",
        domain_status,
        tr("domain_name"),
    )
)

private_button.clicked.connect(
    lambda: turn_on_firewall_profile(
        "privateprofile",
        private_status,
        tr("private_name"),
    )
)

public_button.clicked.connect(
    lambda: turn_on_firewall_profile(
        "publicprofile",
        public_status,
        tr("public_name"),
    )
)

firewall_layout = QVBoxLayout()

for widget in (
    firewall_status,
    domain_status,
    domain_button,
    private_status,
    private_button,
    public_status,
    public_button,
):
    firewall_layout.addWidget(widget)

firewall_tab.setLayout(firewall_layout)


# ============================================================
 GERENCIADOR DE SENHAS
# ============================================================

if password_vault.has_master_password():
    password_status = QLabel(tr("password_locked"))
else:
    password_status = QLabel(tr("password_create_master"))

master_password_input = QLineEdit()
master_password_input.setPlaceholderText(
    tr("password_master_placeholder")
)
master_password_input.setEchoMode(QLineEdit.Password)

unlock_password_button = QPushButton(
    tr("password_unlock_button")
)

service_input = QLineEdit()
service_input.setPlaceholderText(
    tr("password_service_placeholder")
)

username_input = QLineEdit()
username_input.setPlaceholderText(
    tr("password_username_placeholder")
)

saved_password_input = QLineEdit()
saved_password_input.setPlaceholderText(
    tr("password_password_placeholder")
)
saved_password_input.setEchoMode(QLineEdit.Password)

save_password_button = QPushButton(
    tr("password_save_button")
)

password_table = QTableWidget()
password_table.setColumnCount(4)
password_table.setHorizontalHeaderLabels([
    tr("password_table_id"),
    tr("password_table_service"),
    tr("password_table_username"),
    tr("password_table_password"),
])
password_table.setColumnHidden(0, True)

show_password_button = QPushButton(
    tr("password_show_button")
)
delete_password_button = QPushButton(
    tr("password_delete_button")
)

selected_password_status = QLabel(
    tr("selected_hidden")
)


def selected_password_id():
    row = password_table.currentRow()

    if row < 0:
        return None

    item = password_table.item(row, 0)

    if item is None:
        return None

    return int(item.text())


def refresh_password_table():
    password_cache.clear()
    password_table.setRowCount(0)

    if not password_vault.is_unlocked():
        return

    try:
        entries = password_vault.list_passwords()

    except PasswordVaultError as error:
        password_status.setText(
            tr("password_error", error=error)
        )
        return

    for row_index, entry in enumerate(entries):
        password_table.insertRow(row_index)

        password_table.setItem(
            row_index,
            0,
            QTableWidgetItem(str(entry["id"])),
        )

        password_table.setItem(
            row_index,
            1,
            QTableWidgetItem(entry["service"]),
        )

        password_table.setItem(
            row_index,
            2,
            QTableWidgetItem(entry["username"]),
        )

        password_table.setItem(
            row_index,
            3,
            QTableWidgetItem("********"),
        )

        password_cache[entry["id"]] = entry["password"]


def unlock_password_manager():
    try:
        if password_vault.has_master_password():
            password_vault.unlock(
                master_password_input.text()
            )
            password_status.setText(
                tr("password_unlocked")
            )

        else:
            password_vault.create_master_password(
                master_password_input.text()
            )
            password_status.setText(
                tr("password_master_created")
            )

        master_password_input.clear()
        selected_password_status.setText(
            tr("selected_hidden")
        )

        refresh_password_table()

    except PasswordVaultError as error:
        password_status.setText(
            tr("password_error", error=error)
        )


def save_password_entry():
    try:
        password_vault.add_password(
            service_input.text(),
            username_input.text(),
            saved_password_input.text(),
        )

        service_input.clear()
        username_input.clear()
        saved_password_input.clear()

        selected_password_status.setText(
            tr("selected_hidden")
        )

        password_status.setText(
            tr("password_saved")
        )

        refresh_password_table()

    except PasswordVaultError as error:
        password_status.setText(
            tr("password_error", error=error)
        )


def show_selected_password():
    password_id = selected_password_id()

    if password_id is None:
        selected_password_status.setText(
            tr("selected_select_first")
        )
        return

    selected_password_status.setText(
        tr(
            "selected_value",
            password=password_cache.get(
                password_id,
                tr("selected_not_found"),
            ),
        )
    )


def delete_selected_password():
    password_id = selected_password_id()

    if password_id is None:
        password_status.setText(
            tr("password_select_first")
        )
        return

    try:
        password_vault.delete_password(password_id)

        selected_password_status.setText(
            tr("selected_hidden")
        )

        password_status.setText(
            tr("password_deleted")
        )

        refresh_password_table()

    except PasswordVaultError as error:
        password_status.setText(
            tr("password_error", error=error)
        )


unlock_password_button.clicked.connect(
    unlock_password_manager
)
save_password_button.clicked.connect(
    save_password_entry
)
show_password_button.clicked.connect(
    show_selected_password
)
delete_password_button.clicked.connect(
    delete_selected_password
)

password_layout = QVBoxLayout()

for widget in (
    password_status,
    master_password_input,
    unlock_password_button,
    service_input,
    username_input,
    saved_password_input,
    save_password_button,
    password_table,
    show_password_button,
    delete_password_button,
    selected_password_status,
):
    password_layout.addWidget(widget)

password_tab.setLayout(password_layout)


home_status = QLabel(tr("home_secure"))
home_scan = QLabel(tr("home_scan_default"))
home_firewall = QLabel(tr("home_firewall_default"))

home_layout = QVBoxLayout()
home_layout.addWidget(home_status)
home_layout.addWidget(home_scan)
home_layout.addWidget(home_firewall)

home_tab.setLayout(home_layout)


# ============================================================
# IDIOMAS
# ============================================================

language_status = QLabel(tr("language_status"))
language_description = QLabel(tr("language_description"))

ptbr_button = QPushButton(tr("language_ptbr"))
enuk_button = QPushButton(tr("language_enuk"))


def update_ui_language():
    if last_scan_result:
        scanned_count, found_count = last_scan_result

        status_label.setText(
            tr(
                "status_scan_completed",
                scanned=scanned_count,
                found=found_count,
            )
        )

        home_scan.setText(
            tr(
                "home_scan_result",
                scanned=scanned_count,
                found=found_count,
            )
        )

    else:
        status_label.setText(tr("status_secure"))
        home_scan.setText(tr("home_scan_default"))

    scan_button.setText(tr("quick_scan"))
    custom_scan_button.setText(tr("custom_scan"))

    if protection_needs_action:
        home_status.setText(tr("home_action_needed"))
    else:
        home_status.setText(tr("home_secure"))

    if home_firewall_enabled:
        home_firewall.setText(tr("home_firewall_enabled"))
    else:
        home_firewall.setText(tr("home_firewall_default"))

    firewall_status.setText(tr("firewall_title"))
    domain_status.setText(tr("domain_unknown"))
    private_status.setText(tr("private_unknown"))
    public_status.setText(tr("public_unknown"))

    domain_button.setText(tr("domain_button"))
    private_button.setText(tr("private_button"))
    public_button.setText(tr("public_button"))

    if password_vault.is_unlocked():
        password_status.setText(tr("password_unlocked"))
    elif password_vault.has_master_password():
        password_status.setText(tr("password_locked"))
    else:
        password_status.setText(tr("password_create_master"))

    master_password_input.setPlaceholderText(
        tr("password_master_placeholder")
    )
    unlock_password_button.setText(
        tr("password_unlock_button")
    )
    service_input.setPlaceholderText(
        tr("password_service_placeholder")
    )
    username_input.setPlaceholderText(
        tr("password_username_placeholder")
    )
    saved_password_input.setPlaceholderText(
        tr("password_password_placeholder")
    )

    save_password_button.setText(
        tr("password_save_button")
    )

    password_table.setHorizontalHeaderLabels([
        tr("password_table_id"),
        tr("password_table_service"),
        tr("password_table_username"),
        tr("password_table_password"),
    ])

    show_password_button.setText(
        tr("password_show_button")
    )
    delete_password_button.setText(
        tr("password_delete_button")
    )

    selected_password_status.setText(
        tr("selected_hidden")
    )

    language_status.setText(tr("language_status"))
    language_description.setText(
        tr("language_description")
    )
    ptbr_button.setText(tr("language_ptbr"))
    enuk_button.setText(tr("language_enuk"))

    tabs.setTabText(0, tr("tab_home"))
    tabs.setTabText(1, tr("tab_scan"))
    tabs.setTabText(2, tr("tab_firewall"))
    tabs.setTabText(3, tr("tab_password"))
    tabs.setTabText(4, tr("tab_language"))


def set_language(language):
    global current_language

    current_language = language
    update_ui_language()


ptbr_button.clicked.connect(
    lambda: set_language("pt_br")
)
enuk_button.clicked.connect(
    lambda: set_language("en_uk")
)

language_layout = QVBoxLayout()
language_layout.addWidget(language_status)
language_layout.addWidget(language_description)
language_layout.addWidget(ptbr_button)
language_layout.addWidget(enuk_button)

language_tab.setLayout(language_layout)
 

tabs.addTab(home_tab, tr("tab_home"))
tabs.addTab(scan_tab, tr("tab_scan"))
tabs.addTab(firewall_tab, tr("tab_firewall"))
tabs.addTab(password_tab, tr("tab_password"))
tabs.addTab(language_tab, tr("tab_language"))

scan_layout = QVBoxLayout()
scan_layout.addWidget(scan_button)
scan_layout.addWidget(custom_scan_button)
scan_layout.addWidget(
    QLabel(tr("preview_hint"))
)
scan_layout.addWidget(result_list)

scan_tab.setLayout(scan_layout)

layout = QVBoxLayout()
layout.addWidget(title_label)
layout.addWidget(status_label)
layout.addWidget(tabs)

window.setLayout(layout)
window.show()

sys.exit(app.exec())
